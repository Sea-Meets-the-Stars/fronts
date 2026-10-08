"""Build the M2 acceptance deck.

Content is quoted from the M2 task logs in ``claude_prompts/frontogenesis_prompts.md``
("Execution prompt 3", tasks 1-6 and the M2-Q7 entry; 2026-10-03 .. 2026-10-04) and from
the Status, criteria and Q&A of ``claude_prompts/frontogenesis_prompt_3.md``.  Where a later
entry corrects an earlier number the later one is used (task 6 corrects task 5's "so
``xr.merge`` works").  Figures come from ``make_m2_figs.py`` (run it first).  Nothing is
recomputed and no data store is opened.

House rule, as for M1: **no text smaller than 20 pt** -- every run, including tags, captions
and footers.  ``check_m1_deck.py Frontogenesis_M2_Acceptance.pptx`` verifies that after the
build.

Style (palette, helpers) is copied from ``build_m1_deck.py`` rather than imported, since
importing it would run the M1 build.

Run:  <frontogenesis env python> make_m2_figs.py && <...> build_m2_deck.py
Deps: python-pptx (in the frontogenesis env since M0).
"""
import pathlib
from pptx import Presentation
from pptx.util import Inches as I, Pt
from pptx.dml.color import RGBColor as C
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

HERE = pathlib.Path(__file__).parent
FIGS = HERE / "figs_m2"

MIDNIGHT, DEEP, TEAL = C(0x21,0x29,0x5C), C(0x06,0x5A,0x82), C(0x1C,0x72,0x93)
MINT, WHITE, INK = C(0x8F,0xBF,0xD0), C(0xFF,0xFF,0xFF), C(0x1B,0x2A,0x33)
MUTED, LIGHT, CARD = C(0x62,0x73,0x7F), C(0xF2,0xF6,0xF8), C(0xE7,0xEF,0xF3)
AMBER, AMBERBG = C(0xA8,0x50,0x1B), C(0xFA,0xEF,0xE4)
GREEN = C(0x2C,0x6E,0x49)
HEAD, BODY, MONO = "Cambria", "Calibri", "Courier New"
MIN_PT = 20          # the house rule

prs = Presentation(); prs.slide_width, prs.slide_height = I(13.333), I(7.5)
BLANK = prs.slide_layouts[6]


def slide(bg=LIGHT):
    s = prs.slides.add_slide(BLANK)
    r = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    r.fill.solid(); r.fill.fore_color.rgb = bg
    r.line.fill.background(); r.shadow.inherit = False
    return s


def _run(p, t, size, bold, font, italic, color):
    r = p.add_run(); r.text = t; f = r.font
    f.size = Pt(max(size, MIN_PT)); f.bold = bold; f.name = font
    f.italic = italic; f.color.rgb = color
    return r


def txt(s, x, y, w, h, runs, size=20, color=INK, font=BODY, bold=False,
        align=PP_ALIGN.LEFT, space=4, line=1.1, italic=False):
    """Text box.  ``runs`` is a str, or a list of paragraphs; a paragraph is a str or a
    (text, opts) tuple, and ``text`` may itself be a list of (str, opts) runs."""
    tb = s.shapes.add_textbox(I(x), I(y), I(w), I(h)); tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    if isinstance(runs, str):
        runs = [(runs, {})]
    for i, item in enumerate(runs):
        t, o = item if isinstance(item, tuple) else (item, {})
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = o.get("align", align)
        p.space_after = Pt(o.get("space", space))
        p.line_spacing = o.get("line", line)
        pieces = t if isinstance(t, list) else [(t, {})]
        for k, (pt, po) in enumerate(pieces):
            if k == 0 and o.get("bullet"):
                pt = "•  " + pt
            oo = {**o, **po}
            _run(p, pt, oo.get("size", size), oo.get("bold", bold), oo.get("font", font),
                 oo.get("italic", italic), oo.get("color", color))
    return tb


def card(s, x, y, w, h, fill=CARD):
    sh = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, I(x), I(y), I(w), I(h))
    sh.fill.solid(); sh.fill.fore_color.rgb = fill
    sh.line.fill.background(); sh.adjustments[0] = 0.07; sh.shadow.inherit = False
    return sh


def circle(s, x, y, d, label, fill=DEEP, fg=WHITE, size=20):
    sh = s.shapes.add_shape(MSO_SHAPE.OVAL, I(x), I(y), I(d), I(d))
    sh.fill.solid(); sh.fill.fore_color.rgb = fill
    sh.line.fill.background(); sh.shadow.inherit = False
    tf = sh.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.word_wrap = False          # "M2" / "Q7" must not break into two lines
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    _run(p, label, size, True, BODY, False, fg)
    return sh


def title(s, text, sub=None, dark=False, tag=None):
    y = 0.38
    if tag:
        txt(s, 0.85, y, 11.7, 0.36, tag, size=20, bold=True,
            color=MINT if dark else TEAL, space=0)
        y += 0.36
    txt(s, 0.85, y, 11.7, 0.65, text, size=30, bold=True, font=HEAD,
        color=WHITE if dark else MIDNIGHT, space=0)
    if sub:
        txt(s, 0.85, y + 0.66, 11.7, 0.42, sub, size=20,
            color=MINT if dark else MUTED, space=0)


def pic(s, name, x, y, w=None, h=None):
    p = FIGS / name
    if p.exists():
        kw = {"width": I(w)} if w else {"height": I(h)}
        s.shapes.add_picture(str(p), I(x), I(y), **kw)
    return p.exists()


def bullets(s, x, y, w, h, items, size=20, space=8, color=INK):
    return txt(s, x, y, w, h, [(t, {"bullet": True}) for t in items],
               size=size, color=color, space=space)


def heading(s, x, y, w, text, color=DEEP):
    return txt(s, x, y, w, 0.4, text, size=20, bold=True, color=color, space=0)


def caption(s, x, y, w, text):
    return txt(s, x, y, w, 0.4, text, size=20, font=MONO, color=MUTED, space=0)


# ------------------------------------------------------------------ 1 TITLE
s = slide(MIDNIGHT)
for d, cx, cy, col in ((4.4, 10.2, 3.1, DEEP), (3.0, 11.2, 4.6, TEAL), (1.8, 10.4, 5.7, MINT)):
    o = s.shapes.add_shape(MSO_SHAPE.OVAL, I(cx - d / 2), I(cy - d / 2), I(d), I(d))
    o.fill.background(); o.line.color.rgb = col; o.line.width = Pt(1.4)
    o.shadow.inherit = False
txt(s, 0.95, 1.3, 8.4, 0.4, "MILESTONE M2  ·  DATA PULL  ·  TWO SOURCES, ONE WINDOW",
    size=20, color=MINT, bold=True, space=0)
txt(s, 0.95, 1.95, 8.6, 1.9,
    [("Frontogenesis in the", {}), ("LLC4320 California Current", {})],
    size=38, bold=True, font=HEAD, color=WHITE, space=4, line=1.05)
txt(s, 0.95, 4.05, 8.2, 1.0,
    "The 72-hour window on disk, one time-dimensioned store per source — resumably, "
    "and verified hour by hour.",
    size=20, color=MINT, line=1.2)
card(s, 0.95, 5.2, 8.2, 0.95, DEEP)
txt(s, 1.3, 5.42, 7.6, 0.5,
    "M2 closed 2026-10-04 — all seven criteria PASS; 72 / 72 hours in both stores",
    size=21, bold=True, color=WHITE, space=0)
txt(s, 0.95, 6.5, 8.2, 0.5,
    "J. Xavier Prochaska  ·  tile 330, face 10  ·  2026-10-03 / 04",
    size=20, color=MUTED, space=0)

# ------------------------------------------------------------------ 2 CONTENTS
s = slide()
title(s, "Contents", "One summary, six tasks and one Q&A finding, the audit, and what M3 inherits")
left = [
    ("M2", "The answer: M2 in one slide", TEAL),
    ("1", "pull_series and verify_series", DEEP),
    ("2", "The 72-hour OSN pull", DEEP),
    ("3", "QA of the series", DEEP),
    ("3+", "The V3 baseline across the window", DEEP),
    ("Q7", "Tile-edge margin test", AMBER),
]
right = [
    ("4", "Chunk-store reconnaissance", DEEP),
    ("5", "load_chunk_levels and the chunk pull", DEEP),
    ("6", "M2 acceptance — the audit", GREEN),
    ("→", "Carried to M3", TEAL),
    ("§", "Glossary", TEAL),
    ("7", "This deck", DEEP),
]
for col, x in ((left, 0.9), (right, 7.0)):
    for i, (n, h, c) in enumerate(col):
        y = 1.9 + i * 0.72
        circle(s, x, y, 0.5, n, c, size=20)
        txt(s, x + 0.7, y + 0.08, 5.3, 0.4, h, size=20, bold=True, color=MIDNIGHT, space=0)

# ------------------------------------------------------------------ 3 M2 IN ONE SLIDE
# The executive view: numbers from the task-2, 4, 5, 6 logs and the M2-Q7 entry.
s = slide()
title(s, "The answer: M2 in one slide",
      "Goal: the 72-hour window on disk, one time-dimensioned store per source, resumably",
      tag="MILESTONE M2  ·  DATA PULL  ·  CLOSED 2026-10-04")
stats = [
    ("72 / 72 hours", "in both stores: OSN surface, 765 MiB in 27.3 min; chunk k = 0..2, "
                      "973 MiB in 7.03 h"),
    ("bit-identical", "chunk k = 0 Theta / Salt / W and Eta equal OSN's in all 72 hours — "
                      "one model output"),
    ("0 mismatches", "land-NaN vs hFacC / W / S in every hour and level; both re-runs a "
                     "no-op, every file sha256-identical"),
]
for i, (big, lab) in enumerate(stats):
    x = 0.85 + i * 3.95
    card(s, x, 1.85, 3.75, 1.8)
    txt(s, x + 0.2, 1.92, 3.35, 0.5, big, size=24, bold=True, font=HEAD, color=DEEP, space=0)
    txt(s, x + 0.2, 2.42, 3.35, 1.2, lab, size=20, color=INK, line=1.0, space=0)
heading(s, 0.85, 3.78, 5.6, "WHAT WAS FOUND")
bullets(s, 0.85, 4.18, 5.6, 2.2, [
    "Flux attrs say +=down but the data are upward-positive: negated at write, stored "
    "downward-positive (M2-Q6 a).",
    "V3 gate: 64 of 71 pairs pass (0.972 ± 0.020) — a real front, not the tile edge; "
    "edge_cells stays 7 (M2-Q7).",
], space=3)
heading(s, 6.9, 3.78, 5.6, "HANDED TO M3")
bullets(s, 6.9, 4.18, 5.6, 2.2, [
    "Two stores sharing time, niter, j, i, XC, YC; a plain xr.merge fails — rename or "
    "subset first.",
    "W.isel(k_l=1); drF 1.0 / 1.14 / 1.30 m; 6-hourly forcing; displacement max 2.27 "
    "cells on the mask.",
], space=3)
card(s, 0.85, 6.5, 11.65, 0.65, MIDNIGHT)
txt(s, 1.2, 6.65, 11.0, 0.45,
    "All seven criteria PASS — 127 passed, 3 xfailed, smoke tests pass. M2 closed 2026-10-04.",
    size=20, bold=True, color=WHITE, space=0)

# ------------------------------------------------------------------ 4 TASK 1
s = slide()
title(s, "pull_series and verify_series",
      "Resumable and atomic per hour; 'present' is the store's own time coord, never a side file",
      tag="TASK 1  ·  CRITERION 2 IN CODE, PROVEN ON REAL DATA IN TASKS 2 AND 5")
card(s, 0.85, 1.9, 5.7, 4.35)
heading(s, 1.15, 2.05, 5.1, "ATOMICITY — MEASURED, THEN BUILT")
bullets(s, 1.15, 2.5, 5.1, 3.7, [
    "to_zarr(append_dim='time') is not atomic: a kill after the third array left time at "
    "length 2 and Salt / Eta / niter at 1 — the half-hour the prompt warns about.",
    "So: stage the hour in memory and write once; append only the time-dimensioned "
    "variables; repair on resume — truncate to the shortest array, drop trailing slabs "
    "that are unwritten, unreadable or all fill.",
], space=6)
card(s, 6.75, 1.9, 5.75, 4.35, AMBERBG)
heading(s, 7.05, 2.05, 5.2, "STOP AT THE FIRST GAP", color=AMBER)
bullets(s, 7.05, 2.5, 5.2, 3.7, [
    "Retries: 3 attempts, backoff 5 / 20 / 60 s. An hour that still fails is recorded and "
    "the run returns — criterion 1 is 'no gaps', and an append-along-time store cannot "
    "take a middle hour later.",
    "The next run retries that hour first: an interrupted pull is restarted, not debugged.",
], space=6)
txt(s, 0.85, 6.45, 11.65, 0.7,
    "zarr_series.py (247 lines) is generic, so task 5 reuses it. 20 offline tests + 1 "
    "network smoke test (22 s for one real hour); suite 104 passed + 3 xfailed.",
    size=20, color=MUTED, line=1.05, space=0)

# ------------------------------------------------------------------ 5 TASK 2
s = slide()
title(s, "The 72-hour OSN pull",
      "data/tile330_raw_20120702T00_72h.zarr — launched detached, finished on the first launch",
      tag="TASK 2  ·  CRITERIA 1, 2, 3 AND THE OSN HALF OF 7")
pic(s, "m2_fig_osn_walltime.png", 0.85, 1.95, w=5.6)
caption(s, 0.85, 5.35, 5.6, "data/m2_pull_done_run1.json")
txt(s, 0.85, 5.85, 5.6, 1.3,
    "OSN was uniformly fast: a 21-33 s band, none of the 90-192 s stalls M0 saw. Network-bound; "
    "writing an hour is ~0.2 s.",
    size=20, color=MUTED, line=1.05, space=0)
bullets(s, 6.8, 1.95, 5.7, 5.1, [
    "72 hours in 27.3 min: median 22.4 s per hour, range 20.7-33.1 s; 0 retries, 0 failures, "
    "0 repairs.",
    "765 MiB on disk (11.1 MB/hour, 834 files). verify_series OK in 0.9 s: no gaps, §3.2 "
    "schema, land-NaN = hFac in all 72 hours, niter steps 144, KPPhbl present.",
    "No-op proof: re-run 0.9 s, 0 pulled / 72 skipped; sha256 + stat of all 834 files, "
    "0 differing lines.",
    "Hours 0-1 identical to M0's 2-hour store, which is kept (M2-Q4). Scale-up: the 504-hour "
    "series ~3.2 h, 5.6 GB.",
], space=6)

# ------------------------------------------------------------------ 6 TASK 3
s = slide()
title(s, "QA of the series",
      "Tide, diurnal mixed layer, constant land-NaN, and the displacement envelope for M3",
      tag="TASK 3  ·  figs/m2_qa_series.png, RE-PLOTTED FROM ITS CACHED SUMMARIES")
pic(s, "m2_fig_qa_series.png", 1.55, 1.9, w=10.2)
bullets(s, 0.85, 4.6, 11.65, 2.5, [
    "Eta range 2.009 m, M2 0.619 + K1 0.494 m (r² 0.989). KPPhbl amplitude 6.4 m, maximum at "
    "~01 h local solar; the afternoon minimum deepens 13 → 12 → 8 m as |tau| falls "
    "0.11 → 0.06 N m⁻².",
    "Land-NaN constant, 0 mismatches vs hFac in every hour; oceTAU* 922 / 565 finite on land → "
    "0 after re-masking; no frozen field, NaN change or outlier.",
    "Displacement: ocean median 0.27-0.44, p99 1.05-1.38, max 4.05 cells (Gulf of California "
    "jet, outside the mask); mask_analysis max 2.27. At L = 8, 7-34 cells per pair lose "
    "support: edge_cells = 7 kept, isfinite required.",
], space=4)

# ------------------------------------------------------------------ 7 TASK 3 EXTRA (M2-Q3)
s = slide()
title(s, "The V3 baseline across the window",
      "0.981 is a property of the operators, with a leverage-driven tail on day 3",
      tag="TASK 3, EXTRA STEP (M2-Q3)  ·  figs/m2_v3_stability.png")
pic(s, "m2_fig_stability.png", 0.85, 1.9, w=6.9)
caption(s, 0.85, 6.25, 6.9, "m2_v3_stability.json; fails: 36, 62-67")
bullets(s, 8.0, 1.95, 4.5, 5.1, [
    "Pair 0-1 reproduces M1: 0.9806 [0.9698, 0.9942].",
    "All 71 pairs: 0.902-0.997, mean 0.972 ± 0.020, weighted 0.983; 64/71 pass; every CI "
    "overlaps the baseline band.",
    "Top 1 % |2F| trimmed: 0.971-1.002, 0.987 ± 0.005 — leverage from one sharp front, not "
    "a drift of the kinematics.",
    "Figure 2 keeps 0.981 [0.970, 0.994]; M3 quotes the window spread as the temporal "
    "systematic.",
], space=6)

# ------------------------------------------------------------------ 8 M2-Q7
# Table and verdict from the M2-Q7 log entry (2026-10-03).
s = slide()
title(s, "Tile-edge margin test",
      "Hypothesis: the day-3 dips are tile-edge contamination. Verdict: a real 2.2 °C front; "
      "edge_cells stays 7",
      tag="M2-Q7  ·  JXP: OPTION (a), TEST IT NOW  ·  figs/m2_q7_edge_margin.png")
cols = [("edge_cells", 0.85, 1.3), ("pass", 2.25, 0.95), ("failing pairs", 3.2, 1.55),
        ("mean ± std", 4.75, 1.85), ("trimmed", 6.65, 1.85)]
rows = [
    ("7", "64/71", "36, 62-67", "0.972 ± 0.020", "0.987 ± 0.005"),
    ("10", "69/71", "36, 63", "0.979 ± 0.013", "0.988 ± 0.005"),
    ("13", "69/71", "36, 70", "0.980 ± 0.012", "0.987 ± 0.005"),
    ("16", "68/71", "36, 69, 70", "0.978 ± 0.013", "0.986 ± 0.005"),
]
y0 = 1.95
for name, x, w in cols:
    txt(s, x, y0, w, 0.4, name, size=20, bold=True, color=DEEP, space=0)
for r, vals in enumerate(rows):
    y = y0 + 0.47 * (r + 1)
    if r == 0:
        card(s, 0.75, y - 0.04, 7.8, 0.44, CARD)
    for (name, x, w), v in zip(cols, vals):
        txt(s, x, y, w, 0.4, v, size=20, bold=(r == 0), color=INK, space=0)
bullets(s, 0.85, 4.45, 7.7, 2.7, [
    "Crop test: moving the edge 4-16 cells inward changes values only within 3-5 cells of the "
    "new edge (as NaN); beyond that, ≤ 4e-12 — round-off.",
    "A wider margin clears 62-67 only by excluding the front; 36 fails at every width, and "
    "69-70 start failing at 13-16 from interior fronts.",
    "Raw Theta: 11.84 → 14.02 °C across i = 7-11 at hour 64. The 13-cell row goes to M3 as a "
    "sensitivity.",
], space=5)
pic(s, "m2_crop_q7c.png", 8.8, 1.9, h=4.5)
caption(s, 8.8, 6.5, 3.9, "panel (c), j = 176")

# ------------------------------------------------------------------ 9 TASK 4
s = slide()
title(s, "Chunk-store reconnaissance",
      "s3://dbof/LLC4320_RAW/CHUNKS/monterey_bay/ on Nautilus: 72 / 72 hours, exactly tile 330",
      tag="TASK 4  ·  SOURCE B  ·  CRITERION 6 AND THE HOUR INVENTORY FOR 5")
pic(s, "m2_fig_flux_diurnal.png", 0.85, 1.95, w=5.6)
caption(s, 0.85, 5.15, 5.6, "data/m2_chunk_recon.json, 07-03")
txt(s, 0.85, 5.65, 5.6, 1.5,
    "Per hour on the store: Theta 57.2, Salt 47.5, U 63.8, V 64.2, W 65.4 MB, each 2-D field "
    "~1.2 MB — 306 MB compressed, 22 GB for 72 hours.",
    size=20, color=MUTED, line=1.05, space=0)
bullets(s, 6.8, 1.95, 5.7, 5.1, [
    "One zarr v3 store per hour, 51 levels, oceQsw and oceFWflx present; grid bit-identical "
    "to tile330_grid.zarr. W sits on k_p1 (52): k_p1 = n is the top face of cell n.",
    "Not level-selective: one 51-level zstd object per variable per hour, so k = 0..2 fetches "
    "174 MB/hour, 12.5 GB; at 0.55 MB/s here, ~6.3 h.",
    "drF[0] = 1.0 m and Z[0] = −0.5 m confirmed; drF 1.0 / 1.14 / 1.30 m. k = 0 fields and "
    "W(k_p1=0) bit-identical to OSN; Eta identical in all 72 hours.",
    "Flux attrs say +=down, the data are upward-positive (oceQsw ≤ 0, −589 W/m² at 13 LST); "
    "the forcing is 6-hourly, linearly interpolated.",
], space=5)

# ------------------------------------------------------------------ 10 TASK 5
s = slide()
title(s, "load_chunk_levels and the chunk pull",
      "data/tile330_chunk_20120702T00_72h.zarr — §3.3; fluxes negated to downward-positive "
      "(M2-Q6 a)",
      tag="TASK 5  ·  CRITERIA 5, 6 AND THE CHUNK HALF OF 7")
pic(s, "m2_fig_chunk_walltime.png", 0.85, 1.95, w=5.6)
caption(s, 0.85, 5.6, 5.6, "data/m2_chunk_pull_done_run1.json")
txt(s, 0.85, 6.05, 5.6, 1.1,
    "First launch stopped at hour 0 by a Salt bound (0, 45): the Gulf of California reaches "
    "48.6 psu. Bound widened; stop-at-gap proven on real data.",
    size=20, color=MUTED, line=1.05, space=0)
bullets(s, 6.8, 1.95, 5.7, 5.1, [
    "72 / 72 hours, 0 missing, in 7.03 h: median 299 s per hour, range 297-1887 s; 12.6 GB "
    "fetched, 973 MiB on disk.",
    "Two FSTimeoutErrors on W, recovered on attempt 2; two slow reads; one 28-min stall from "
    "idle sleep — caffeinate -i -s next time.",
    "Every object validated before any write: decode, size, shape, NaN = (hFacC == 0), "
    "plausibility, time / iteration attrs, chunk Eta bit-identical to OSN.",
    "verify_chunk_series OK: 0 land-NaN mismatches over 72 h × 3 levels, oceQsw ≥ 0. Re-run "
    "a no-op, all 692 files sha256-identical. 23 offline tests + 1 network.",
], space=4)

# ------------------------------------------------------------------ 11 TASK 6: THE AUDIT
s = slide(MIDNIGHT)
title(s, "M2 acceptance — the audit",
      "All seven criteria, as audited in the task-6 log entry; the 'Do not' list checked on disk",
      dark=True, tag="TASK 6")
crit = [
    ("72 steps, no gaps, schema", "both stores 72 / 72, steps 3600 s; both verifies re-run "
                                  "fresh, OK"),
    ("No-op re-run", "0 pulled, 72 skipped; sha256 of 834 + 692 files identical"),
    ("KPPhbl present", "every hour, finite on every ocean cell; 6.4 m diurnal"),
    ("Land-NaN = hFac", "0 mismatch cells, 72 hours × 3 levels: M1's masks hold"),
    ("Chunk k = 0..2, fluxes, drF", "72 / 72 hours, missing: none; downward-positive"),
    ("drF[0] confirmed", "1.0 m (then 1.14, 1.30); Z[0] = −0.5 m"),
    ("Volume and wall time", "OSN 765 MiB, 27.3 min; chunk 973 MiB, 7.03 h, 12.6 GB fetched"),
]
y = 2.0
for h, body in crit:
    circle(s, 0.9, y + 0.03, 0.42, "✓", GREEN, size=20)
    txt(s, 1.55, y, 10.9, 0.7,
        [([(h + "  ", {"bold": True, "color": WHITE}), (body, {"color": MINT})], {})],
        size=20, space=0, line=1.05)
    y += 0.6
card(s, 0.85, 6.4, 11.6, 0.75, DEEP)
txt(s, 1.25, 6.57, 10.8, 0.45,
    "127 passed, 3 strict xfailed; both smoke tests pass. M2 closed 2026-10-04.",
    size=20, bold=True, color=WHITE, space=0)

# ------------------------------------------------------------------ 12 CARRIED TO M3
# From the task-6 "Carried forward to M3" list (items 1-8, 10).
s = slide()
title(s, "Carried to M3",
      "The task-6 audit's carried-forward list, for frontogenesis_prompt_4.md",
      tag="FOR M3  ·  THE RULES THE DATA IMPOSE")
heading(s, 0.85, 1.95, 5.6, "THE DATA")
bullets(s, 0.85, 2.4, 5.6, 4.8, [
    "Merge: a plain xr.merge of the two stores fails on Theta / Salt / W (2-D vs 3-D). "
    "Rename them (16 vars) or subset to the fluxes, drF and W.isel(k_l=1) (14 vars).",
    "The fluxes are already downward-positive: surface_flux_term must not negate again.",
    "W.isel(k_l=1) is the cell-base velocity; W(k_l=0) = dEta/dt is a free-surface signal.",
    "Forcing is 6-hourly, linearly interpolated: the shortwave diurnal is a triangle peaking "
    "at 13 LST. Phase Figure 6 on the mixed-layer minimum.",
], space=5)
heading(s, 6.9, 1.95, 5.6, "THE STATISTICS")
bullets(s, 6.9, 2.4, 5.6, 4.8, [
    "isfinite(DGDt) & mask_analysis at L = 8; edge_cells = 7 kept; the 13-cell row "
    "(0.980 ± 0.012) as a sensitivity.",
    "Trimmed and orthogonal fits beside the OLS gate for every pair; 0.972 ± 0.020 as the "
    "temporal systematic; show 07-04 14-20 UTC separately, not dropped.",
    "Displacement envelope: analysis max 2.27 cells; ocean 4.05 in the excluded Gulf.",
    "Pending: the note to Lauren on the flux attrs, for JXP to forward; caffeinate -i -s "
    "for long pulls on the laptop.",
], space=5)

# ------------------------------------------------------------------ 13 GLOSSARY
s = slide()
title(s, "Glossary — the M2-specific terms", tag="GLOSSARY  ·  NOTATION AS ON THE SLIDES")
txt(s, 0.85, 1.55, 11.65, 5.6,
    [([(term + "  ", {"bold": True, "color": DEEP}), (defn, {})], {"space": 5})
     for term, defn in [
        ("Stop-at-gap",
         "the pull records the first hour that still fails after 3 retries and returns, so "
         "the store stays a contiguous prefix of the series; the next run retries that hour."),
        ("Repair-on-resume",
         "before trusting the store's time coord, repair_trailing truncates every "
         "time-dimensioned array to the shortest and drops trailing slabs that are unwritten, "
         "unreadable or all fill."),
        ("No-op re-run",
         "a second run pulls nothing and leaves every file sha256-identical — criterion 2, "
         "proven on both stores."),
        ("k_p1, k_l",
         "the source puts W on 52 interfaces k_p1; k_p1 = n is the top face of cell n, i.e. "
         "k_l = n, renamed at write. W(k_l=1) is the cell-base velocity."),
        ("niter, mit_iteration",
         "niter is the OSN iteration (steps of 144 per hour); mit_iteration = niter − 10368 "
         "is the chunk source's selected_iteration."),
        ("Sign convention",
         "the source attrs say +=down but the data are upward-positive; the store negates "
         "oceQnet / oceQsw / oceFWflx so that into the ocean is positive, and says so in "
         "sign_convention and source_sign_convention."),
        ("Trimmed slope",
         "the OLS gate recomputed with the top 1 % |2F| front pixels dropped — a leverage "
         "diagnostic (0.987 ± 0.005), not the gate."),
     ]],
    size=20, line=1.0)

# ------------------------------------------------------------------ 14 TASK 7
s = slide()
title(s, "This deck",
      "Reproducible from the logged numbers and the tasks' JSON summaries; nothing below "
      "20 pt; 14 slides",
      tag="TASK 7  ·  2026-10-04")
card(s, 0.85, 1.9, 11.6, 1.5, CARD)
txt(s, 1.25, 2.1, 10.8, 1.2,
    [("Every number here is quoted from the M2 task logs in frontogenesis_prompts.md, from "
      "frontogenesis_prompt_3.md, or read from the tasks' small JSON done-files.", {"bold": True}),
     ("Nothing is recomputed and no zarr store is opened.", {"color": MUTED})],
    size=20, space=4, line=1.05)
files = [
    ("make_m2_figs.py", "one panel crop from ../figs/ and five large-font re-plots"),
    ("build_m2_deck.py", "builds this .pptx"),
    ("check_m1_deck.py", "reused: every run ≥ 20 pt, plus the text-box geometry"),
    ("figs_m2/", "the six images used here"),
    ("README.md", "the work log for this task"),
]
heading(s, 0.85, 3.65, 11.6, "IN dev/frontogenesis/deck/")
for i, (f, d) in enumerate(files):
    y = 4.15 + i * 0.52
    txt(s, 1.1, y, 3.6, 0.4, f, size=20, font=MONO, bold=True, color=MIDNIGHT, space=0)
    txt(s, 4.8, y + 0.02, 7.6, 0.4, d, size=20, color=MUTED, space=0)
txt(s, 0.85, 6.85, 11.6, 0.4,
    "Rendered with LibreOffice and inspected page by page before hand-over.",
    size=20, color=MUTED, space=0)

out = HERE / "Frontogenesis_M2_Acceptance.pptx"
prs.save(out)
print(f"saved {out.name}  ({out.stat().st_size/1024:.0f} KB, {len(prs.slides._sldIdLst)} slides)")
