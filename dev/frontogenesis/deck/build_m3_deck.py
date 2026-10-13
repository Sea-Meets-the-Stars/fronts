"""Build the M3 acceptance deck.

Content is quoted from the M3 task logs in ``claude_prompts/frontogenesis_prompts.md``
("Execution prompt 4", tasks 1-8; 2026-10-07 .. 2026-10-11) and from the Status, criteria,
"Do not" list and Q&A of ``claude_prompts/frontogenesis_prompt_4.md``.  Where a later entry
corrects an earlier number the later one is used (task 7's rebuilt sweep supersedes task 4's
store sizes; task 6's pooled numbers supersede task 3's hour-0 smoke).  Figures come from
``make_m3_figs.py`` (run it first).  **Nothing is recomputed and no data store is opened.**

**M3 did not close.**  The deck says so on the title slide, in the one-slide summary and on
the task-6 slides; no frontogenesis efficiency appears anywhere, and the slopes are shown
only under a "not quoted" label (prompt 4's "Do not" list, first item).

House rule, as for M1 and M2: **no text smaller than 20 pt** -- every run, including tags,
captions and footers.  ``check_m1_deck.py Frontogenesis_M3_Acceptance.pptx`` verifies it.
Figures are placed at their **natural inch size** (they were drawn at slide width with
``font.size = 20``), so their panel labels clear the floor too.

Style (palette, helpers) is copied from ``build_m2_deck.py`` rather than imported, since
importing it would run the M2 build.

Run:  <frontogenesis env python> make_m3_figs.py && <...> build_m3_deck.py
"""
import pathlib
from pptx import Presentation
from pptx.util import Inches as I, Pt
from pptx.dml.color import RGBColor as C
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

HERE = pathlib.Path(__file__).parent
FIGS = HERE / "figs_m3"

MIDNIGHT, DEEP, TEAL = C(0x21,0x29,0x5C), C(0x06,0x5A,0x82), C(0x1C,0x72,0x93)
MINT, WHITE, INK = C(0x8F,0xBF,0xD0), C(0xFF,0xFF,0xFF), C(0x1B,0x2A,0x33)
MUTED, LIGHT, CARD = C(0x62,0x73,0x7F), C(0xF2,0xF6,0xF8), C(0xE7,0xEF,0xF3)
AMBER, AMBERBG = C(0xA8,0x50,0x1B), C(0xFA,0xEF,0xE4)
GREEN, RED, REDBG = C(0x2C,0x6E,0x49), C(0xB2,0x22,0x22), C(0xFA,0xE8,0xE8)
HEAD, BODY, MONO = "Cambria", "Calibri", "Courier New"
MIN_PT = 20

prs = Presentation(); prs.slide_width, prs.slide_height = I(13.333), I(7.5)
BLANK = prs.slide_layouts[6]


def slide(bg=LIGHT):
    s = prs.slides.add_slide(BLANK)
    r = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    r.fill.solid(); r.fill.fore_color.rgb = bg
    r.line.fill.background(); r.shadow.inherit = False
    return s


def _run(p, t, size, bold, font, italic, color):
    r = p.add_run(); r.text = t; f = r.font
    f.size = Pt(max(size, MIN_PT)); f.bold = bold; f.name = font
    f.italic = italic; f.color.rgb = color
    return r


def txt(s, x, y, w, h, runs, size=20, color=INK, font=BODY, bold=False,
        align=PP_ALIGN.LEFT, space=4, line=1.1, italic=False):
    tb = s.shapes.add_textbox(I(x), I(y), I(w), I(h)); tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    if isinstance(runs, str):
        runs = [(runs, {})]
    for i, item in enumerate(runs):
        t, o = item if isinstance(item, tuple) else (item, {})
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = o.get("align", align)
        p.space_after = Pt(o.get("space", space))
        p.line_spacing = o.get("line", line)
        pieces = t if isinstance(t, list) else [(t, {})]
        for k, (pt_, po) in enumerate(pieces):
            if k == 0 and o.get("bullet"):
                pt_ = "•  " + pt_
            oo = {**o, **po}
            _run(p, pt_, oo.get("size", size), oo.get("bold", bold), oo.get("font", font),
                 oo.get("italic", italic), oo.get("color", color))
    return tb


def card(s, x, y, w, h, fill=CARD):
    sh = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, I(x), I(y), I(w), I(h))
    sh.fill.solid(); sh.fill.fore_color.rgb = fill
    sh.line.fill.background(); sh.adjustments[0] = 0.07; sh.shadow.inherit = False
    return sh


def circle(s, x, y, d, label, fill=DEEP, fg=WHITE, size=20):
    sh = s.shapes.add_shape(MSO_SHAPE.OVAL, I(x), I(y), I(d), I(d))
    sh.fill.solid(); sh.fill.fore_color.rgb = fill
    sh.line.fill.background(); sh.shadow.inherit = False
    tf = sh.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.word_wrap = False
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    _run(p, label, size, True, BODY, False, fg)
    return sh


def title(s, text, sub=None, dark=False, tag=None):
    y = 0.34
    if tag:
        txt(s, 0.85, y, 11.7, 0.34, tag, size=20, bold=True,
            color=MINT if dark else TEAL, space=0)
        y += 0.36
    txt(s, 0.85, y, 11.7, 0.62, text, size=30, bold=True, font=HEAD,
        color=WHITE if dark else MIDNIGHT, space=0)
    if sub:
        txt(s, 0.85, y + 0.64, 11.7, 0.42, sub, size=20,
            color=MINT if dark else MUTED, space=0)


def pic(s, name, x, y, w):
    p = FIGS / name
    if p.exists():
        s.shapes.add_picture(str(p), I(x), I(y), width=I(w))
    else:
        print(f"  MISSING FIGURE {name}")
    return p.exists()


def bullets(s, x, y, w, h, items, size=20, space=8, color=INK):
    return txt(s, x, y, w, h, [(t, {"bullet": True}) for t in items],
               size=size, color=color, space=space)


def heading(s, x, y, w, text, color=DEEP):
    return txt(s, x, y, w, 0.38, text, size=20, bold=True, color=color, space=0)


#: at 20 pt a 11.7 in box holds about 84 characters on one line (the checker's
#: 0.5 em estimate); a footer longer than that would wrap off the slide
FOOTER_MAX = 84


def footer(s, text):
    assert len(text) <= FOOTER_MAX, f"footer {len(text)} chars > {FOOTER_MAX}: {text}"
    txt(s, 0.85, 6.92, 11.7, 0.36, text, size=20, color=MUTED, space=0)


# ================================================================== 1 TITLE
s = slide(MIDNIGHT)
for d, cx, cy, col in ((4.4, 10.2, 3.1, DEEP), (3.0, 11.2, 4.6, TEAL), (1.8, 10.4, 5.7, MINT)):
    o = s.shapes.add_shape(MSO_SHAPE.OVAL, I(cx - d / 2), I(cy - d / 2), I(d), I(d))
    o.fill.background(); o.line.color.rgb = col; o.line.width = Pt(1.4)
    o.shadow.inherit = False
txt(s, 0.95, 1.25, 8.4, 0.4, "MILESTONE M3  ·  FIELD-LEVEL BUDGET  ·  HARD GATE",
    size=20, color=MINT, bold=True, space=0)
txt(s, 0.95, 1.9, 8.6, 1.9,
    [("Frontogenesis in the", {}), ("LLC4320 California Current", {})],
    size=38, bold=True, font=HEAD, color=WHITE, space=4, line=1.05)
txt(s, 0.95, 4.0, 8.3, 1.1,
    "Can hourly surface fields at 2 km close the buoyancy-gradient budget? "
    "We pre-declared the tolerance, measured every term, and answered.",
    size=20, color=MINT, line=1.2)
card(s, 0.95, 5.25, 8.3, 1.0, RED)
txt(s, 1.3, 5.46, 7.7, 0.55,
    "M3 NOT CLOSED 2026-10-11 — the gate failed; the result is methodological",
    size=21, bold=True, color=WHITE, space=0)
txt(s, 0.95, 6.5, 8.3, 0.5,
    "J. Xavier Prochaska  ·  tile 330, face 10  ·  2026-10-07 / 11",
    size=20, color=MUTED, space=0)

# ================================================================== 2 CONTENTS
s = slide()
title(s, "Contents", "One summary, eight tasks, the verdict and its diagnosis, and what is carried forward")
left = [("M3", "The answer: M3 in one slide", RED),
        ("1", "inputs.py — the merged input layer", DEEP),
        ("2", "vertical.py — the measured terms", DEEP),
        ("2b", "The workstation environment", DEEP),
        ("3", "budget.py — the budget per pair", DEEP),
        ("4", "The L sweep over 71 pairs", DEEP),
        ("5", "stats.py — estimators and blocks", DEEP)]
right = [("6", "Closure — THE GATE", RED),
         ("6", "Why it failed — the diagnosis", RED),
         ("6", "Criteria 2 and 3, and Figure 2b", AMBER),
         ("7", "The ten figures", DEEP),
         ("8", "M3 acceptance — the audit", AMBER),
         ("→", "Carried to M4 / M5", TEAL),
         ("§", "Glossary  ·  This deck", TEAL)]
for col, x in ((left, 0.9), (right, 7.0)):
    for i, (n, h, c) in enumerate(col):
        y = 1.85 + i * 0.72
        circle(s, x, y, 0.5, n, c, size=20)
        txt(s, x + 0.7, y + 0.07, 5.3, 0.4, h, size=20, bold=True, color=MIDNIGHT, space=0)

# ================================================== 3 M3 IN ONE SLIDE
s = slide()
title(s, "The answer: M3 in one slide",
      "Goal: close the surface buoyancy-gradient budget per pixel, with every term measured",
      tag="MILESTONE M3  ·  HARD GATE  ·  NOT CLOSED 2026-10-11")
card(s, 0.85, 1.72, 11.7, 1.06, CARD)
txt(s, 1.1, 1.88, 11.2, 0.8,
    "measured  −  2F  −  subfilter  −  vertical  −  surface flux   ~   numerical + interior KPP",
    size=22, bold=True, font=MONO, color=MIDNIGHT, space=0, align=PP_ALIGN.CENTER)
stats = [("0.81 – 1.04", "residual / measured", RED),
         ("0 of 71", "pairs pass, any L", RED),
         ("2, 3, 5", "criteria that PASS", GREEN),
         ("no", "efficiency quoted", AMBER)]
for i, (big, sub, col) in enumerate(stats):
    x = 0.85 + i * 2.97
    card(s, x, 2.92, 2.75, 1.36, WHITE)
    txt(s, x + 0.12, 3.04, 2.51, 0.5, big, size=26, bold=True, color=col, space=0,
        align=PP_ALIGN.CENTER)
    txt(s, x + 0.12, 3.62, 2.51, 0.4, sub, size=20, color=INK, space=0,
        align=PP_ALIGN.CENTER)
card(s, 0.85, 4.45, 11.7, 2.3, REDBG)
txt(s, 1.15, 4.62, 11.1, 0.42,
    "The verdict, against tolerances fixed by JXP on 2026-10-07 and never changed",
    size=20, bold=True, color=RED, space=2)
bullets(s, 1.15, 5.08, 11.1, 1.6, [
    "Criterion 1 fails at every judged L. The residual is as large as the measured tendency.",
    "Two separable causes: at L = 0–2 the hourly sampling; at L = 4–8 the surface-flux term.",
    "Planning §12's null result — hourly 2 km fields cannot constrain this budget. Nothing tuned."],
    size=20, space=5)
footer(s, "Tolerance (M3-Q1, M3-Q9): residual/measured ≤ 0.50, judged at L ≥ 2")

# ================================================================== 4 TASK 1
s = slide()
title(s, "inputs.py — one door to the data", "The only place M3 opens a store", tag="TASK 1")
bullets(s, 0.85, 1.75, 11.7, 2.4, [
    "A plain xr.merge of the two stores fails (Salt is 2-D in one and 3-D in the other), so they merge by RENAME: Theta_k, Salt_k, W_k — 16 variables (M3-Q6).",
    "Two invariants asserted on every hour a budget is computed from: the chunk store's k = 0 is bit-identical to OSN, and the fluxes carry sign_convention = positive downward with oceQsw ≥ 0.",
    "Accessors keep the physics out of the store: W_k1 (never k_l = 0, which is dEta/dt), drF, Z, fluxes, wind, hour_pair, midpoint, filtered, valid."], size=20, space=10)
card(s, 0.85, 4.35, 11.7, 1.5, CARD)
heading(s, 1.1, 4.52, 11.2, "The V3 / M2 pool reproduces exactly on the real stores")
txt(s, 1.1, 4.92, 11.2, 0.85,
    "262,925 valid cells and 26,293 front pixels at L = 0;  drF = [1.0, 1.14, 1.30];  "
    "the k = 0 identity holds on all 72 hours.",
    size=20, color=INK, space=0, line=1.15)
footer(s, "py/inputs.py + tests/test_inputs.py (14 tests) — log entry 2026-10-07, task 1")

# ================================================================== 5 TASK 2
s = slide()
title(s, "vertical.py — the terms that make it a real budget",
      "Measured from the chunk store, not assumed", tag="TASK 2")
bullets(s, 0.85, 1.7, 11.7, 2.3, [
    "b_z from the chunk levels; the vertical term formed as a TENDENCY first, then differenced — the factorised form drops −w grad(b_z)·grad b (M3-Q11).",
    "The surface-flux term from oceQnet, oceQsw and oceFWflx with JMD95 alpha/beta, the shortwave split out: f_sw = 0.521 in the 1 m cell (Jerlov IA, M3-Q10), rhoConst = 1027.5 (M3-Q12).",
    "Smoke at L = 0: the vertical term is 2.7–6.5 % of 2F, the surface-flux term 0.65–0.94 of 2F."], size=20, space=10)
card(s, 0.85, 4.2, 11.7, 1.75, AMBERBG)
heading(s, 1.1, 4.38, 11.2, "A finding, not a tuning", AMBER)
txt(s, 1.1, 4.78, 11.2, 1.05,
    "The surface-flux term is dominated by the NON-SOLAR flux gradient and does not peak at "
    "13 LST. The warm layer is 0.01 K, not the 0.1–0.3 K planning §2.2 assumed, so the vertical "
    "term has almost no diurnal contrast.",
    size=20, color=INK, space=0, line=1.15)
footer(s, "py/vertical.py + py/chunk_store.py + tests/test_vertical.py (19 tests), task 2")

# ================================================================== 6 TASK 2b
s = slide()
title(s, "The workstation environment", "M3 moved from the laptop to profx", tag="TASK 2b")
rows = [("Platform", "Ubuntu 24.04, x86_64, 24 threads, 755 GB RAM — the laptop was macOS arm64"),
        ("Environment", "conda-forge, Python 3.13.16; EVERY pinned version matches the laptop exactly"),
        ("Dropped", "torch, torchvision, healpy, pyvista, PyQt6, timm — none is imported by M3"),
        ("Verified", "suite 160 passed + 3 xfails in 194 s; both 72-hour stores verify ok"),
        ("Long jobs", "nohup inside tmux — caffeinate is macOS-only and this box never suspends")]
for i, (k, v) in enumerate(rows):
    y = 1.8 + i * 0.95
    card(s, 0.85, y, 11.7, 0.82, WHITE if i % 2 == 0 else CARD)
    txt(s, 1.1, y + 0.1, 2.5, 0.4, k, size=20, bold=True, color=DEEP, space=0)
    txt(s, 3.7, y + 0.1, 8.6, 0.62, v, size=20, color=INK, space=0, line=1.1)
footer(s, "Logged: importing every module ran three unguarded scripts; content verified")

# ================================================================== 7 TASK 3
s = slide()
title(s, "budget.py — one hour pair, one filter scale",
      "Every term, and a report that reads like a verdict", tag="TASK 3")
bullets(s, 0.85, 1.7, 11.7, 2.2, [
    "compute_budget returns the §3.4 fields on (time, j, i): b, G, 2F in both forms, the semi-Lagrangian and Eulerian tendencies, the subfilter, vertical and surface-flux terms, the residual and the diagnostics.",
    "No silent degradation: without the chunk store the two measured terms are ABSENT, not zero, and closure_report returns closed = None with a warning in its printed text.",
    "L = 0 is labelled: the explicit subfilter term is identically zero there, not small."], size=20, space=10)
card(s, 0.85, 4.1, 11.7, 1.9, CARD)
heading(s, 1.1, 4.28, 11.2, "Three bit-for-bit identities on the real hour 0")
txt(s, 1.1, 4.68, 11.2, 1.15,
    "2F IS validate.two_F(form='discrete');  DGDt_semilag IS validate.null_step's measured with "
    "the real b(t+1) substituted;  and the pool is 262,925 valid / 26,293 front. "
    "Same functions, same inputs — the plumbing adds nothing.",
    size=20, color=INK, space=0, line=1.15)
footer(s, "py/budget.py + tests/test_budget.py (32 tests) — log entry 2026-10-10, task 3")

# ================================================================== 8 TASK 4
s = slide()
title(s, "The L sweep — 284 budgets over 71 pairs",
      "L ∈ {0, 2, 4, 8}, resumable, verified", tag="TASK 4")
pic(s, "m3_fig_sweep.png", 1.07, 1.75, 11.2)
card(s, 0.85, 6.02, 11.7, 0.8, CARD)
txt(s, 1.1, 6.14, 11.2, 0.6,
    [("71 pairs at every L · no gaps · n_lost 0 at L ≤ 4, max 31 at L = 8 (bound 34)", {}),
     ("The re-run is a no-op: all 7,596 files sha256-unchanged", {})],
    size=20, color=INK, space=0, line=1.1)
footer(s, "M3-Q3's optional L = 1 is impossible: lowpass refuses an odd L (half-width L/2)")

# ================================================================== 9 TASK 5
s = slide()
title(s, "stats.py — estimators, blocks and the baseline",
      "Planning §11's four mechanisms, each pinned by a test", tag="TASK 5")
rows = [("Attenuation", "noise on x attenuates OLS to b·var(X)/(var(X)+var(η)); TLS does not — and the bisector OVERSHOOTS, so it is not a fix either"),
        ("Leverage", "one heavy-tailed point moves OLS 0.897 → 0.760 and the trimmed estimate not at all"),
        ("Ratio by sign", "on a two-signed pool the unsplit ratio reads 1.18 where OLS reads 1.000 — at L = 8 on real data it reads 0.088"),
        ("Blocks, not pixels", "the block bootstrap is 15× wider than a pixel bootstrap; 3-hour blocks are wider again ONLY when the series is correlated")]
for i, (k, v) in enumerate(rows):
    y = 1.8 + i * 1.12
    card(s, 0.85, y, 11.7, 1.0, WHITE if i % 2 == 0 else CARD)
    txt(s, 1.1, y + 0.1, 2.9, 0.4, k, size=20, bold=True, color=DEEP, space=0)
    txt(s, 4.1, y + 0.1, 8.2, 0.82, v, size=20, color=INK, space=0, line=1.1)
footer(s, "Every slope is divided by the M1 baseline 0.981 — never by 1 (coding §8)")

# ================================================================== 10 TASK 6 VERDICT
s = slide()
title(s, "Closure — the hard gate", "The five-term budget, pooled over all 71 pairs, on front & valid",
      tag="TASK 6  ·  CRITERION 1  ·  FAIL AT EVERY JUDGED L")
pic(s, "m3_fig_budget.png", 1.65, 1.82, 10.04)
footer(s, "Pre-declared 2026-10-07: ≤ 0.50. Measured 0.813 / 0.841 / 0.898 / 1.042")

# ================================================================== 11 TASK 6 PER PAIR
s = slide()
title(s, "Not one pair of seventy-one", "The residual fraction for every pair, at every L",
      tag="TASK 6  ·  CRITERION 1")
pic(s, "m3_fig_per_pair.png", 1.65, 1.68, 10.04)
card(s, 0.85, 6.3, 11.7, 0.52, CARD)
txt(s, 1.1, 6.39, 11.2, 0.38,
    "The day-3 pairs 62–68 are BETTER than the median — M2's anomaly does not drive this.",
    size=20, color=INK, space=0)
footer(s, "Per-pair median 0.831 / 0.861 / 0.910 / 1.048, sd 0.07–0.10, full range 0.715–1.245")

# ================================================================== 12 TASK 6 DIAGNOSIS A
s = slide()
title(s, "Why — the terms do not explain the residual",
      "+1 for a correct term, −1 for a sign error, ~0 for one that is simply irrelevant",
      tag="TASK 6  ·  THE DIAGNOSIS")
pic(s, "m3_fig_term_contribution.png", 1.78, 1.88, 9.78)
footer(s, "Near-orthogonal, NOT sign-flipped: a sign error would sit at −1, not at +0.05")

# ================================================================== 13 TASK 6 DIAGNOSIS B
s = slide()
title(s, "Why — and what the hourly cadence costs",
      "Two estimates of the SAME DG/Dt, and the interpolation-order control",
      tag="TASK 6  ·  THE DIAGNOSIS")
pic(s, "m3_fig_sampling.png", 1.65, 1.82, 10.04)
footer(s, "M3-Q14 (a): only hourly data exists, so this cannot be tested at a shorter dt")

# ================================================================== 14 TASK 6 FIG 2b
s = slide()
title(s, "The discriminator", "Numerical diffusion, or air-sea forcing? Neither, at any scale",
      tag="TASK 6  ·  FIGURE 2b")
pic(s, "m3_fig_discriminator.png", 1.56, 1.70, 10.21)
footer(s, "'Do not': no damping may be claimed without Figure 2b — and it refutes the reading")

# ================================================================== 15 CRITERIA 2 AND 3
s = slide()
title(s, "Criteria 2 and 3 — both pass", "The two supporting criteria, over all 71 pairs",
      tag="TASK 6  ·  CRITERIA 2, 3")
pic(s, "m3_fig_criteria23.png", 1.07, 1.82, 11.2)
card(s, 0.85, 6.3, 11.7, 0.5, CARD)
txt(s, 1.1, 6.38, 11.2, 0.36,
    "Criterion 2 passes at L ≥ 4; criterion 3 reproduces M1 task 4 over the whole window.",
    size=20, color=INK, space=0)
footer(s, "Criterion 2 at L = 0 reproduces M1 task 3's slope 0.73 / corr 0.74 almost exactly")

# ================================================================== 16 TASK 7
s = slide()
title(s, "The ten figures", "Built from the stores and the task-6 summary only — no physics recomputed",
      tag="TASK 7  ·  CRITERION 5")
names = [("1", "measured, 2F and residual as maps"), ("2", "the joint PDF — SLOPES NOT QUOTED"),
         ("2b", "the discriminator (above)"), ("3", "slope, criterion 2 and the subfilter vs L"),
         ("3b", "the filter ladder; L = 0 tau = 0"), ("4", "gradient–strain alignment"),
         ("5", "sharpening time and the 1 h contour"), ("6", "the diurnal composite"),
         ("7", "statistics vs distance offshore"), ("10", "the five-term budget")]
for i, (n, h) in enumerate(names):
    col, row = i // 5, i % 5
    x, y = 0.9 + col * 6.1, 1.8 + row * 0.78
    circle(s, x, y, 0.5, n, DEEP, size=20)
    txt(s, x + 0.68, y + 0.07, 5.2, 0.4, h, size=20, color=MIDNIGHT, space=0)
card(s, 0.85, 5.9, 11.7, 1.0, AMBERBG)
txt(s, 1.1, 6.05, 11.2, 0.75,
    "M3-Q13 (c) applied first: the surface-flux term also computed over KPPhbl. It fixes the "
    "term's MAGNITUDE (optimal multiplier +0.05 → +0.6–2.0) and changes nothing about closure.",
    size=20, color=INK, space=0, line=1.1)
footer(s, "One prompt expectation is unmet: the residual does NOT change sign over the day")

# ================================================================== 17 TASK 8
s = slide()
title(s, "M3 acceptance — the audit", "Every criterion, with numbers", tag="TASK 8")
hdr = [("#", 0.9, 0.5), ("criterion", 1.5, 3.5), ("measured", 5.1, 4.9), ("verdict", 10.1, 2.4)]
for t, x, w in hdr:
    txt(s, x, 1.72, w, 0.38, t, size=20, bold=True, color=DEEP, space=0)
rows = [("1", "Closure ≤ 0.50", "0.813 / 0.841 / 0.898 / 1.042", "FAIL", RED),
        ("2", "Eulerian vs semi-Lag.", "slope 0.891 / 0.939 at L = 4 / 8", "PASS at L ≥ 4", GREEN),
        ("3", "Filter sweep interpretable", "subfilter/2F 0.353 / 0.552 / 0.748", "PASS", GREEN),
        ("4", "Slopes vs the baseline", "machinery complete; slopes NOT quoted", "PASS*", AMBER),
        ("5", "Ten figures written", "10 / 10, tracked in git", "PASS", GREEN)]
for i, (n, c, m, v, col) in enumerate(rows):
    y = 2.15 + i * 0.78
    card(s, 0.85, y, 11.7, 0.68, WHITE if i % 2 == 0 else CARD)
    txt(s, 1.0, y + 0.12, 0.5, 0.4, n, size=20, bold=True, color=MIDNIGHT, space=0)
    txt(s, 1.5, y + 0.12, 3.5, 0.4, c, size=20, color=INK, space=0)
    txt(s, 5.1, y + 0.12, 4.9, 0.4, m, size=20, color=INK, space=0)
    txt(s, 10.1, y + 0.12, 2.4, 0.4, v, size=20, bold=True, color=col, space=0)
card(s, 0.85, 6.1, 11.7, 0.82, REDBG)
txt(s, 1.1, 6.25, 11.2, 0.55,
    "Criterion 1 is the gate. M3 is NOT closed; by coding §7 everything downstream stays blocked.",
    size=20, bold=True, color=RED, space=0)
footer(s, "*machinery passes; the slopes may not be quoted while criterion 1 fails")

# ================================================================== 18 CARRIED
s = slide()
title(s, "Carried to M4 / M5", "And what blocks first", tag="→")
card(s, 0.85, 1.7, 5.75, 4.55, CARD)
heading(s, 1.1, 1.9, 5.2, "To M4 — which does not start")
bullets(s, 1.1, 2.38, 5.25, 3.7, [
    "The four derived stores and the per-pair closure JSONs",
    "The slopes — which may NOT be quoted as an efficiency",
    "Three fronts bugs and their strict xfails",
    "Prompt 5 still names a non-existent build.tile_find path"], size=20, space=5)
card(s, 6.8, 1.7, 5.75, 4.55, CARD)
heading(s, 7.05, 1.9, 5.2, "To M5")
bullets(s, 7.05, 2.38, 5.25, 3.7, [
    "The four standard captions, for standardisation",
    "The one-command harness, deliberately not built here",
    "Five module sizes over the ~400-line guidance, with cut points",
    "The V-figures, until §4.10 is consolidated"], size=20, space=5)
footer(s, "M3-Q13 (c), M3-Q14 (a), M3-Q15 (a) answered 2026-10-10 and applied")

# ================================================================== 19 OPEN QUESTIONS
s = slide()
title(s, "Open questions for JXP", "Nothing in the Q&A is unanswered; these are decisions, not physics",
      tag="→")
items = [("Does M4 proceed on a null M3?",
          "Coding §7 blocks everything after M3 on closure. The derived stores, the per-pair "
          "closure JSONs and the slopes are ready for it either way."),
         ("The note to Lauren on flux signs",
          "Written 2026-10-07, committed, and still awaiting forwarding. No M3 action was "
          "possible; it is recorded so it is not lost."),
         ("Is a longer window ever wanted?",
          "The budget is not the bottleneck: 3.6 h and 53 GB for 504 hours, against the chunk "
          "pull's 42 h. A server-side k = 0..2 subset would cut that ~7x.")]
for i, (k, v) in enumerate(items):
    y = 1.8 + i * 1.72
    card(s, 0.85, y, 11.7, 1.55, AMBERBG if i == 0 else CARD)
    heading(s, 1.1, y + 0.14, 11.2, k, AMBER if i == 0 else DEEP)
    txt(s, 1.1, y + 0.58, 11.2, 0.85, v, size=20, color=INK, space=0, line=1.12)
footer(s, "Five module sizes over coding §1.3's guidance are recorded for M5, with cut points")

# ================================================================== 20 GLOSSARY
s = slide()
title(s, "Glossary", "The terms this deck uses without explaining them")
items = [("G", "|grad b|² — front strength, whose budget this is"),
         ("2F", "twice the frontogenesis tendency; F = ½ DG/Dt"),
         ("measured", "the semi-Lagrangian DG/Dt over one hour"),
         ("subfilter", "what the filter hides; zero by construction at L = 0"),
         ("L", "filter scale in cells — the same top-hat on b, U and V"),
         ("front & valid", "G above its p90 at the midpoint, inside mask_analysis"),
         ("baseline", "the M1 discrete null 0.981 — every slope is quoted against it"),
         ("lap²b", "the grad⁴-like diagnostic: numerical diffusion would track it"),
         ("KPPhbl", "the boundary-layer depth: air-sea forcing would track it")]
for i, (k, v) in enumerate(items):
    y = 1.78 + i * 0.58
    txt(s, 0.9, y, 2.3, 0.4, k, size=20, bold=True, font=MONO, color=DEEP, space=0)
    txt(s, 3.3, y, 9.2, 0.45, v, size=20, color=INK, space=0)
footer(s, "Front OBJECTS do not exist in M3 — no labelling or thinning; that is M4")

# ================================================================== 21 THIS DECK
s = slide()
title(s, "This deck", "Built by script from the M3 logs; nothing recomputed")
bullets(s, 0.85, 1.75, 11.7, 2.7, [
    "deck/make_m3_figs.py — seven re-plots at slide size from the task-6 and task-7 JSON summaries. Re-plots, not crops: the page figures' labels fall below 20 pt on a slide.",
    "deck/build_m3_deck.py — this file. Every number is quoted from the task-1..8 log entries; no data store is opened.",
    "check_m1_deck.py verifies the 20 pt floor and the geometry; rendered with soffice + pdftoppm and every page inspected."], size=20, space=9)
card(s, 0.85, 4.62, 11.7, 2.2, REDBG)
heading(s, 1.1, 4.78, 11.2, "What this deck does not contain", RED)
bullets(s, 1.1, 5.16, 11.2, 1.6, [
    "No frontogenesis efficiency. Criterion 1 failed, so the 'Do not' list forbids one.",
    "No claim of diabatic damping. Figure 2b is the discriminator and it refutes the reading.",
    "No tuned number. The tolerances were fixed on 2026-10-07 and are asserted in the test suite."], size=20, space=5)
footer(s, "Frontogenesis_M3_Acceptance.pptx — 21 slides, 2026-10-11")

out = HERE / "Frontogenesis_M3_Acceptance.pptx"
prs.save(out)
print(f"wrote {out}  ({out.stat().st_size/1024:.0f} KB, {len(prs.slides._sldIdLst)} slides)")
