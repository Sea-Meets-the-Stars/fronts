from pptx import Presentation
from pptx.util import Inches as I, Pt
from pptx.dml.color import RGBColor as C
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

MIDNIGHT, DEEP, TEAL = C(0x21,0x29,0x5C), C(0x06,0x5A,0x82), C(0x1C,0x72,0x93)
MINT, WHITE, INK = C(0x8F,0xBF,0xD0), C(0xFF,0xFF,0xFF), C(0x1B,0x2A,0x33)
MUTED, LIGHT, CARD = C(0x62,0x73,0x7F), C(0xF2,0xF6,0xF8), C(0xE7,0xEF,0xF3)
AMBER, AMBERBG = C(0xA8,0x50,0x1B), C(0xFA,0xEF,0xE4)
HEAD, BODY, MONO = "Cambria", "Calibri", "Courier New"

prs = Presentation(); prs.slide_width, prs.slide_height = I(13.333), I(7.5)
BLANK = prs.slide_layouts[6]

def slide(bg=LIGHT):
    s = prs.slides.add_slide(BLANK)
    r = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    r.fill.solid(); r.fill.fore_color.rgb = bg; r.line.fill.background(); r.shadow.inherit = False
    return s

def txt(s, x, y, w, h, runs, size=15, color=INK, font=BODY, bold=False,
        align=PP_ALIGN.LEFT, space=6, line=None, italic=False):
    tb = s.shapes.add_textbox(I(x), I(y), I(w), I(h)); tf = tb.text_frame
    tf.word_wrap = True; tf.margin_left = tf.margin_right = 0
    tf.margin_top = tf.margin_bottom = 0
    if isinstance(runs, str): runs = [(runs, {})]
    for i, item in enumerate(runs):
        t, o = item if isinstance(item, tuple) else (item, {})
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = o.get("align", align)
        p.space_after = Pt(o.get("space", space))
        if line or o.get("line"): p.line_spacing = o.get("line", line)
        if o.get("bullet"): t = "•   " + t
        r = p.add_run(); r.text = t
        f = r.font
        f.size = Pt(o.get("size", size)); f.bold = o.get("bold", bold)
        f.name = o.get("font", font); f.italic = o.get("italic", italic)
        f.color.rgb = o.get("color", color)
    return tb

def card(s, x, y, w, h, fill=CARD, radius=True, shadow=False):
    sh = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
                            I(x), I(y), I(w), I(h))
    sh.fill.solid(); sh.fill.fore_color.rgb = fill; sh.line.fill.background()
    if radius: sh.adjustments[0] = 0.06
    sh.shadow.inherit = False
    return sh

def circle(s, x, y, d, label, fill=DEEP, fg=WHITE, size=15):
    sh = s.shapes.add_shape(MSO_SHAPE.OVAL, I(x), I(y), I(d), I(d))
    sh.fill.solid(); sh.fill.fore_color.rgb = fill; sh.line.fill.background()
    sh.shadow.inherit = False
    tf = sh.text_frame; tf.margin_left = tf.margin_right = 0
    tf.margin_top = tf.margin_bottom = 0; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = label
    r.font.size = Pt(size); r.font.bold = True; r.font.name = BODY; r.font.color.rgb = fg
    return sh

def title(s, text, sub=None, dark=False):
    txt(s, 0.85, 0.52, 11.7, 0.85, text, size=34, bold=True, font=HEAD,
        color=WHITE if dark else MIDNIGHT)
    if sub:
        txt(s, 0.85, 1.36, 11.7, 0.42, sub, size=14,
            color=MINT if dark else MUTED, space=0)

# ---------------------------------------------------------------- 1 TITLE
s = slide(MIDNIGHT)
for d, cx, cy, col in ((4.6, 9.9, 3.0, DEEP), (3.2, 11.0, 4.5, TEAL), (1.9, 10.2, 5.6, MINT)):
    o = s.shapes.add_shape(MSO_SHAPE.OVAL, I(cx - d/2), I(cy - d/2), I(d), I(d))
    o.fill.background(); o.line.color.rgb = col; o.line.width = Pt(1.5); o.shadow.inherit = False
txt(s, 0.95, 1.55, 8.4, 0.35, "LLC4320  ·  CALIFORNIA CURRENT  ·  SURFACE",
    size=12, color=MINT, bold=True, space=0)
txt(s, 0.95, 2.15, 8.6, 2.3,
    [("Frontogenesis in the", {}), ("LLC4320 California Current", {})],
    size=40, bold=True, font=HEAD, color=WHITE, space=4, line=1.05)
txt(s, 0.95, 4.35, 8.2, 0.9,
    "Comparing measured front strengthening against the frontogenesis rate "
    "computed from the model's own velocity and buoyancy fields.",
    size=15, color=MINT, line=1.3)
txt(s, 0.95, 6.25, 8.2, 0.6,
    [("Planning Summary", {"bold": True, "color": WHITE}),
     ("J. Xavier Prochaska  ·  HIINet  ·  September 2026", {"size": 12, "color": MUTED})],
    size=14, space=2)

# ---------------------------------------------------------------- 2 QUESTION
s = slide()
title(s, "The question", "Two numbers that should agree — and the reason they will not")
card(s, 0.85, 2.05, 5.6, 2.5, CARD)
circle(s, 1.2, 2.35, 0.5, "M", DEEP)
txt(s, 1.9, 2.45, 4.3, 0.4, "MEASURED", size=12, bold=True, color=DEEP, space=0)
txt(s, 1.2, 3.15, 4.9, 1.2,
    "How much fronts actually sharpen, hour to hour, following the flow.",
    size=15, color=INK, line=1.3)
txt(s, 1.2, 3.95, 4.9, 0.4, "D G / D t     from consecutive snapshots",
    size=13, font=MONO, color=MUTED, space=0)
card(s, 6.85, 2.05, 5.6, 2.5, CARD)
circle(s, 7.2, 2.35, 0.5, "P", TEAL)
txt(s, 7.9, 2.45, 4.3, 0.4, "PREDICTED", size=12, bold=True, color=TEAL, space=0)
txt(s, 7.2, 3.15, 4.9, 1.2,
    "How much the strain field says they should sharpen, instantaneously.",
    size=15, color=INK, line=1.3)
txt(s, 7.2, 3.95, 4.9, 0.4, "2F     from u, v and b", size=13, font=MONO, color=MUTED, space=0)
card(s, 0.85, 4.95, 11.6, 1.55, MIDNIGHT)
txt(s, 1.35, 5.3, 10.6, 1.0,
    [("These are not two guesses at one number.", {"bold": True, "color": WHITE}),
     ("They are the two sides of an exact budget — and the residual between them is itself the physical result.",
      {"color": MINT})],
    size=16, space=4, line=1.25)

# ---------------------------------------------------------------- 3 BUDGET
s = slide()
title(s, "One budget, three terms", "The surface buoyancy-gradient budget, stated exactly")
card(s, 0.85, 2.0, 11.6, 1.15, MIDNIGHT)
txt(s, 1.2, 2.32, 10.9, 0.6,
    "D/Dt ( ½ |grad b|² )  =  F   -   b_z (w_x b_x + w_y b_y)   +   grad b · grad B",
    size=17, font=MONO, color=WHITE, space=0, align=PP_ALIGN.CENTER)
cols = [
    ("F", "KINEMATIC", "Already implemented on the native grid, with correct dxC/dyC metrics and CS/SN rotation.", DEEP),
    ("T", "TILTING", "Vanishes identically at the surface: w → 0 there, so w_x and w_y do too.", TEAL),
    ("B", "DIABATIC", "Everything we cannot compute directly. This is the residual — and the result.", AMBER),
]
for i, (sym, head, bodytext, col) in enumerate(cols):
    x = 0.85 + i * 3.93
    card(s, x, 3.45, 3.7, 2.05, AMBERBG if col == AMBER else CARD)
    circle(s, x + 0.32, 3.75, 0.46, sym, col)
    txt(s, x + 0.95, 3.85, 2.6, 0.35, head, size=11.5, bold=True, color=col, space=0)
    txt(s, x + 0.32, 4.5, 3.1, 1.3, bodytext, size=13, color=INK, line=1.25)
txt(s, 0.85, 5.85, 11.6, 0.5,
    "Convention: F = ½ D G/Dt, so every comparison is 2F against the measured tendency.",
    size=13, bold=True, color=MIDNIGHT, align=PP_ALIGN.CENTER, space=0)

# ---------------------------------------------------------------- 4 SURFACE-ONLY
s = slide()
title(s, "Why surface-only is enough", "and the one place the discretisation takes it back")
card(s, 0.85, 2.0, 5.6, 3.15, CARD)
circle(s, 1.2, 2.3, 0.5, "✓", DEEP, size=17)
txt(s, 1.9, 2.42, 4.2, 0.4, "IN THE CONTINUUM", size=12, bold=True, color=DEEP, space=0)
txt(s, 1.2, 3.1, 4.9, 1.9,
    "w vanishes identically along the surface, so its along-surface derivatives vanish too. "
    "The tilting term drops out exactly — there is no missing kinematic term.",
    size=14.5, color=INK, line=1.3)
txt(s, 1.2, 4.55, 4.9, 0.4, "Surface data is sufficient, not a compromise.",
    size=13, bold=True, color=DEEP, space=0)
card(s, 6.85, 2.0, 5.6, 3.15, AMBERBG)
circle(s, 7.2, 2.3, 0.5, "!", AMBER, size=17)
txt(s, 7.9, 2.42, 4.2, 0.4, "IN THE DATA", size=12, bold=True, color=AMBER, space=0)
txt(s, 7.2, 3.1, 4.9, 1.9,
    "k = 0 is a finite cell (~1 m), not the mathematical surface. Its budget carries the flux "
    "through the cell base, where w does not vanish.",
    size=14.5, color=INK, line=1.3)
txt(s, 7.2, 4.55, 4.9, 0.4, "~30% of F by day. ~0 at night.",
    size=13, bold=True, color=AMBER, space=0)
txt(s, 0.85, 5.55, 11.6, 0.95,
    "Bounded in Phase 2 against the full-depth monterey_bay chunk store, which has 11 snapshots "
    "inside our window — rather than deferred and hoped about.",
    size=13.5, color=MUTED, align=PP_ALIGN.CENTER, line=1.3)

# ---------------------------------------------------------------- 5 EXISTS/NEW
s = slide()
title(s, "Half the study is already built", "The predicted side exists, and is built well")
have = ["Frontogenesis tendency F on the native grid",
        "Metric-correct gradients and velocity Jacobian",
        "Front detection (NaN-safe)",
        "Front tracking across time",
        "Per-front property statistics",
        "Front-following cross-sections (curtains)"]
new = ["Measured D G/Dt from consecutive snapshots",
       "Semi-Lagrangian departure-point machinery",
       "Coarse-graining and subfilter flux",
       "Concat-to-zarr for a time series",
       "Discrete null-test harness",
       "Cross-front delta-b, width, peak"]
card(s, 0.85, 2.05, 5.6, 4.3, CARD)
txt(s, 1.25, 2.4, 4.9, 0.4, "ALREADY EXISTS", size=12, bold=True, color=DEEP, space=0)
txt(s, 1.25, 3.0, 4.9, 3.1, [(t, {"bullet": True}) for t in have],
    size=13.5, color=INK, space=11, line=1.15)
card(s, 6.85, 2.05, 5.6, 4.3, MIDNIGHT)
txt(s, 7.25, 2.4, 4.9, 0.4, "THE NEW WORK", size=12, bold=True, color=MINT, space=0)
txt(s, 7.25, 3.0, 4.9, 3.1, [(t, {"bullet": True}) for t in new],
    size=13.5, color=WHITE, space=11, line=1.15)
txt(s, 0.85, 6.6, 11.6, 0.4,
    "Nobody has yet measured how fast the fronts actually sharpen. That is the whole gap.",
    size=13, italic=True, color=MUTED, align=PP_ALIGN.CENTER, space=0)

# ---------------------------------------------------------------- 6 DATA
s = slide(MIDNIGHT)
title(s, "The data", "Tile 330 — the best-supported path in the repository", dark=True)
stats = [("720²", "native cells"), ("72", "hourly steps"), ("~1.9 km", "grid spacing"),
         ("0.9 GB", "total volume")]
for i, (big, lab) in enumerate(stats):
    x = 0.85 + i * 2.95
    card(s, x, 2.1, 2.7, 1.5, DEEP)
    txt(s, x + 0.2, 2.35, 2.3, 0.6, big, size=27, bold=True, font=HEAD,
        color=WHITE, align=PP_ALIGN.CENTER, space=0)
    txt(s, x + 0.2, 3.05, 2.3, 0.35, lab, size=11.5, color=MINT,
        align=PP_ALIGN.CENTER, space=0)
txt(s, 0.85, 4.0, 5.6, 2.4,
    [("SOURCE", {"size": 11.5, "bold": True, "color": MINT, "space": 5}),
     ("The public OSN store (anonymous, hourly, surface). Face 10; box ≈ 128–113°W, 26.7–38.2°N. "
      "A second OSN store supplies KPPhbl and wind stress over the same window.",
      {"size": 13.5, "color": WHITE, "line": 1.3})])
txt(s, 6.85, 4.0, 5.6, 2.4,
    [("WINDOW  ·  2 – 4 JULY 2012", {"size": 11.5, "bold": True, "color": MINT, "space": 5}),
     ("Chosen, not defaulted: the full-depth chunk store holds a dense 3-hourly day on 3 July, so this "
      "window captures 11 of its 17 snapshots instead of 3. Free, and it makes the depth cross-check real.",
      {"size": 13.5, "color": WHITE, "line": 1.3})])

# ---------------------------------------------------------------- 7 METHOD
s = slide()
title(s, "Method", "Five commitments that make the comparison mean something")
items = [
    ("One operator, both sides", "Measured and predicted go through the same gradient stencil and the same filter. Different stencils manufacture a mismatch."),
    ("Semi-Lagrangian, not Eulerian", "A single well-conditioned difference along the trajectory, avoiding two large nearly-cancelling terms."),
    ("Identical filtering of b, u and v", "Swept over four scales. Filtering changes the budget, so the sweep is reported as a curve, not a number."),
    ("Land masked before differencing", "Land is stored as zero; untreated, the coastline is the strongest 'front' in the domain."),
    ("Statistics ≥ 100 km offshore", "Cleanest signal, and reported stratified by distance so what we excluded stays visible."),
]
for i, (h, b) in enumerate(items):
    y = 2.0 + i * 0.92
    circle(s, 0.9, y + 0.06, 0.44, str(i + 1), DEEP if i % 2 == 0 else TEAL, size=14)
    txt(s, 1.62, y, 10.8, 0.35, h, size=15, bold=True, color=MIDNIGHT, space=0)
    txt(s, 1.62, y + 0.35, 10.8, 0.5, b, size=12.5, color=MUTED, space=0, line=1.2)

# ---------------------------------------------------------------- 8 HARD PART
s = slide()
title(s, "The residual is not purely diabatic", "The correction that most changes how we read the result")
card(s, 0.85, 2.0, 11.6, 1.25, AMBERBG)
txt(s, 1.3, 2.3, 10.7, 0.7,
    "A regression slope of 0.5–0.8 is fully explicable by the model's own numerics — with zero air-sea flux.",
    size=16.5, bold=True, color=AMBER, align=PP_ALIGN.CENTER, line=1.2)
blocks = [
    ("Numerical diffusion", "LLC4320 has no explicit horizontal tracer diffusion. Its implicit dissipation switches on exactly at the grid-scale gradients that define our fronts — at 0.1–1 f, the same order as the strain."),
    ("Finite-cell vertical term", "Reintroduced by the discretisation even though the continuum tilting term vanishes. ~30% of F by day."),
    ("Discrete operator bias", "The chain rule fails at order unity for 4Δx features, and C-grid staggering attenuates G and F differently — a slope bias of 0.7–1.4 with no physics in it."),
]
for i, (h, b) in enumerate(blocks):
    x = 0.85 + i * 3.93
    card(s, x, 3.5, 3.7, 2.35, CARD)
    txt(s, x + 0.3, 3.8, 3.1, 0.4, h, size=14, bold=True, color=MIDNIGHT, space=0)
    txt(s, x + 0.3, 4.3, 3.1, 1.4, b, size=12, color=INK, line=1.25)
txt(s, 0.85, 6.2, 11.6, 0.45,
    "So “slope < 1 = damping” is not a safe inference. Separating these is now part of the analysis, not a caveat.",
    size=13, bold=True, color=MIDNIGHT, align=PP_ALIGN.CENTER, space=0)

# ---------------------------------------------------------------- 9 GATES
s = slide()
title(s, "Four gates before any science", "Operators must be known correct, not assumed correct")
gates = [
    ("1", "Cartesian scheme test", "Pure deformation, where |grad b|² grows exactly as exp(2αt). Validates the semi-Lagrangian scheme.", DEEP, CARD),
    ("2", "Native-grid metric test", "Analytic function of XC/YC with known gradients. Validates dxC, dyC, CS, SN handling.", TEAL, CARD),
    ("3", "Discrete null test", "Advect a synthetic tracer with our exact discrete operators. Requires slope = 1 ± 0.05 before real data is touched.", AMBER, AMBERBG),
    ("4", "Interpolation-bias test", "Uniform zero-strain flow, where the true tendency is identically zero. Whatever comes out is our bias.", MIDNIGHT, CARD),
]
for i, (n, h, b, col, bg) in enumerate(gates):
    x = 0.85 + (i % 2) * 5.95
    y = 2.05 + (i // 2) * 2.25
    card(s, x, y, 5.6, 2.0, bg)
    circle(s, x + 0.33, y + 0.3, 0.46, n, col)
    txt(s, x + 0.98, y + 0.4, 4.4, 0.35, h, size=14.5, bold=True, color=MIDNIGHT, space=0)
    txt(s, x + 0.33, y + 1.02, 4.95, 0.85, b, size=12.5, color=INK, line=1.25)
txt(s, 0.85, 6.65, 11.6, 0.4,
    "Gate 3 is the one that protects the headline number. Every later slope is quoted against its baseline, never against 1.",
    size=12.5, italic=True, color=MUTED, align=PP_ALIGN.CENTER, space=0)

# ---------------------------------------------------------------- 10 FIGURES
s = slide()
title(s, "What we will show", "Ten figures; these five carry the argument")
figs = [
    ("Maps, side by side", "Measured tendency beside 2F, shared scale, plus residual. If these do not look alike, nothing else matters."),
    ("The joint PDF", "Measured against 2F on front pixels, with the estimators shown against the null-test baseline."),
    ("Numerical vs diabatic", "Residual against high-order derivatives of b and against boundary-layer depth. The discriminator."),
    ("Slope vs filter scale", "Traces how much imbalance is unresolved-scale physics rather than damping."),
    ("Strain alignment", "Angle between grad b and the compressional axis — the classic independent physical check."),
]
for i, (h, b) in enumerate(figs):
    x = 0.85 + (i % 3) * 3.93
    y = 2.05 + (i // 3) * 2.25
    card(s, x, y, 3.7, 2.0, CARD)
    circle(s, x + 0.3, y + 0.28, 0.42, str(i + 1), [DEEP, TEAL, AMBER, DEEP, TEAL][i], size=13)
    txt(s, x + 0.3, y + 0.85, 3.1, 0.35, h, size=13.5, bold=True, color=MIDNIGHT, space=0)
    txt(s, x + 0.3, y + 1.25, 3.1, 0.65, b, size=11.5, color=MUTED, line=1.2)
card(s, 4.78, 4.3, 7.67, 2.0, MIDNIGHT)
txt(s, 5.2, 4.72, 6.9, 1.2,
    [("Plus a tracked-front case study", {"bold": True, "color": WHITE, "size": 15}),
     ("One front followed hour by hour, its observed strength against the F-predicted curve — "
      "and population statistics over the ten largest fronts.", {"color": MINT, "size": 12.5, "line": 1.25})],
    space=5)

# ---------------------------------------------------------------- 11 MILESTONES
s = slide(MIDNIGHT)
title(s, "Milestones", "Each becomes one execution prompt document", dark=True)
ms = [("M0", "Access", False), ("M1", "Operators", True), ("M2", "Data pull", False),
      ("M3", "Budget", True), ("M4", "Fronts", False), ("M5", "Figures", False)]
for i, (m, lab, gate) in enumerate(ms):
    x = 0.85 + i * 2.0
    card(s, x, 2.35, 1.75, 1.5, AMBER if gate else DEEP)
    txt(s, x + 0.1, 2.6, 1.55, 0.45, m, size=22, bold=True, font=HEAD,
        color=WHITE, align=PP_ALIGN.CENTER, space=0)
    txt(s, x + 0.1, 3.18, 1.55, 0.35, lab, size=11.5,
        color=WHITE if gate else MINT, align=PP_ALIGN.CENTER, space=0)
    if gate:
        txt(s, x + 0.1, 3.95, 1.55, 0.3, "GATE", size=10, bold=True,
            color=AMBER, align=PP_ALIGN.CENTER, space=0)
txt(s, 0.85, 4.6, 5.6, 1.9,
    [("M1 — operators", {"size": 13.5, "bold": True, "color": WHITE, "space": 4}),
     ("Does not pass until the discrete null test returns slope = 1 ± 0.05. On failure we fix the "
      "operators, not the interpretation.", {"size": 12.5, "color": MINT, "line": 1.3})])
txt(s, 6.85, 4.6, 5.6, 1.9,
    [("M3 — budget", {"size": 13.5, "bold": True, "color": WHITE, "space": 4}),
     ("Exit criterion is budget closure, not a slope. No efficiency number is quoted before the "
      "terms balance to a stated tolerance.", {"size": 12.5, "color": MINT, "line": 1.3})])
txt(s, 0.85, 6.75, 11.6, 0.35,
    "M1 and M2 can run in parallel — the data pull needs no physics.",
    size=12, italic=True, color=MUTED, align=PP_ALIGN.CENTER, space=0)

# ---------------------------------------------------------------- 12 RISKS
s = slide()
title(s, "One open decision, and the honest risks")
card(s, 0.85, 1.95, 11.6, 1.35, AMBERBG)
circle(s, 1.25, 2.28, 0.5, "?", AMBER, size=18)
txt(s, 2.0, 2.2, 10.1, 0.95,
    [("Blocking: which branches do we work from?", {"bold": True, "size": 15.5, "color": AMBER}),
     ("In fronts, the frontogenesis and viz_tools branches have diverged and both created the same files. "
      "In the preprocessing repo, the checked-out branch is 102 commits behind main.",
      {"size": 12.5, "color": INK, "line": 1.25})], space=4)
risks = [
    ("Implicit numerical diffusion mimics diabatic damping", "at the same order as the strain"),
    ("Discrete operators bias the slope 0.7–1.4", "before any physics enters"),
    ("Semi-Lagrangian interpolation bias", "25–80% of the signal, and signed"),
    ("72 hours cannot separate diurnal, inertial and M2", "3, 3.6 and 5.8 cycles respectively"),
    ("The OSN code path has never been run", "no OSN test, no tile-find test, no NaN-input test"),
]
txt(s, 0.85, 3.55, 11.6, 0.35, "TOP RISKS", size=11.5, bold=True, color=DEEP, space=0)
for i, (h, b) in enumerate(risks):
    y = 4.02 + i * 0.6
    circle(s, 0.9, y + 0.02, 0.28, "", TEAL if i % 2 else DEEP, size=9)
    txt(s, 1.4, y, 6.8, 0.35, h, size=13, bold=True, color=INK, space=0)
    txt(s, 8.3, y + 0.02, 4.15, 0.35, b, size=12, color=MUTED, space=0)

# ---------------------------------------------------------------- 13 NULL
s = slide(MIDNIGHT)
title(s, "What would make this a null result", "Stated now, so we recognise it rather than rationalise around it", dark=True)
nulls = ["The residual is comparable to 2F at all filter scales and the budget does not close.",
         "The measured slope is indistinguishable from the discrete-null baseline.",
         "The residual tracks numerical diffusion rather than boundary-layer depth or the diurnal cycle.",
         "The discrete null test cannot be made to pass at 1 ± 0.05."]
for i, t in enumerate(nulls):
    y = 2.3 + i * 0.78
    circle(s, 0.9, y, 0.42, str(i + 1), DEEP, size=13)
    txt(s, 1.6, y + 0.05, 10.8, 0.5, t, size=14.5, color=WHITE, space=0, line=1.25)
card(s, 0.85, 5.75, 11.6, 1.15, DEEP)
txt(s, 1.3, 6.05, 10.7, 0.65,
    "Then the conclusion is methodological, the limiting factor is named — and that is better "
    "than a slope of 0.6 presented as an efficiency.",
    size=14.5, bold=True, color=WHITE, align=PP_ALIGN.CENTER, line=1.2)

prs.save("Frontogenesis_Planning.pptx")
print("saved", len(prs.slides.__iter__.__self__._sldIdLst), "slides")
